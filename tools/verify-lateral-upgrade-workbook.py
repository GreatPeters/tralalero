"""Read-only scope checks before installing the Artifact Tool-authored workbook."""
from copy import copy
import json
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "Assets/ShooterSurvival/GameData/Editor/Data.xlsx"
CANDIDATE = ROOT / "map-concepts/mobile-presentation-2026-09-14/workbook/Data.xlsx"
before = openpyxl.load_workbook(SOURCE)
after = openpyxl.load_workbook(CANDIDATE)
assert before.sheetnames == after.sheetnames
checked = 0
for sheet in before:
    changed = after[sheet.title]
    assert sheet.freeze_panes == changed.freeze_panes, sheet.title
    assert str(sheet.merged_cells) == str(changed.merged_cells), sheet.title
    assert sheet.sheet_state == changed.sheet_state, sheet.title
    assert sheet.auto_filter.ref == changed.auto_filter.ref, sheet.title
    assert len(sheet.data_validations.dataValidation) == len(changed.data_validations.dataValidation), sheet.title
    assert list(sheet.tables) == list(changed.tables), sheet.title
    for column, dimension in sheet.column_dimensions.items():
        other = changed.column_dimensions[column]
        assert (dimension.width, dimension.hidden, dimension.outlineLevel) == (other.width, other.hidden, other.outlineLevel), (sheet.title, column)
    for row in sheet:
        for cell in row:
            other = changed[cell.coordinate]
            assert cell.value == other.value, (sheet.title, cell.coordinate, cell.value, other.value)
            assert cell.number_format == other.number_format, (sheet.title, cell.coordinate, "number format")
            if cell.has_style:
                for field in ("font", "fill", "border", "alignment", "protection"):
                    assert copy(getattr(cell, field)) == copy(getattr(other, field)), (sheet.title, cell.coordinate, field)
            if cell.comment:
                assert other.comment and cell.comment.text == other.comment.text, (sheet.title, cell.coordinate, "comment")
            if cell.hyperlink:
                assert other.hyperlink and cell.hyperlink.target == other.hyperlink.target, (sheet.title, cell.coordinate, "hyperlink")
            checked += 1
    if sheet.title != "업그레이드":
        assert (sheet.max_row, sheet.max_column) == (changed.max_row, changed.max_column), sheet.title

sheet = after["업그레이드"]
assert sheet.max_row == 312
for level in range(1, 11):
    values = [sheet.cell(302 + level, column).value for column in range(2, 11)]
    assert values[:8] == [10, "LATERAL_SPEED", level, "좌우 이동 속도 늘리기", 5 * level, "percent", "Coin", 35 * (level + 1)], values
    assert sheet.cell(302 + level, 5).alignment.wrap_text
    assert sheet.cell(302 + level, 10).alignment.wrap_text
report = {"unchanged_sheets_and_existing_cells_verified": True, "sheet_count": len(before.sheetnames),
          "existing_cells_checked": checked, "new_rows": 10, "max_level": 10, "max_percent": 50,
          "prices": [35 * (level + 1) for level in range(1, 11)]}
(CANDIDATE.parent / "verification.json").write_text(json.dumps(report, indent=2), encoding="utf-8")
print(json.dumps(report))
